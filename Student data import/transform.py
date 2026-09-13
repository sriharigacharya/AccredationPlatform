import openpyxl
import pandas as pd
from collections import defaultdict, OrderedDict
import re

SRC = "Student_Master_Data.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Master Student Data"]

NROWS = ws.max_row
NCOLS = ws.max_column

# ---- Parse the two-row merged header into course "blocks" ----
header1 = [ws.cell(row=1, column=c).value for c in range(1, NCOLS + 1)]
header2 = [ws.cell(row=2, column=c).value for c in range(1, NCOLS + 1)]

# forward-fill header1 (merged cell titles only appear in the first col of each block)
filled_h1 = []
last = None
for v in header1:
    if v is not None:
        last = v
    filled_h1.append(last)

# course_code / course_name / semester / section extraction from block titles like:
# "Data Structures (CS3C01) - Sem 3 Sec A - Attendance"
# "Computer Org. & Architecture (CS3C03) - Sem 3 Sec B - IA Marks"
BLOCK_RE = re.compile(
    r"^(?P<name>.+?)\s*\((?P<code>[A-Z0-9]+)\)\s*-\s*Sem\s*(?P<sem>\d+)\s*Sec\s*(?P<sec>[A-Z])\s*-\s*(?P<kind>Attendance|IA Marks)$"
)

blocks = OrderedDict()  # key: (col_start) -> dict(meta)
col = 3  # cols 1-2 are USN, Student Name
while col <= NCOLS:
    title = filled_h1[col - 1]
    m = BLOCK_RE.match(title)
    if not m:
        raise ValueError(f"Unrecognized header block at col {col}: {title!r}")
    meta = m.groupdict()
    # gather sub-columns belonging to this block (same title)
    sub_cols = []
    c = col
    while c <= NCOLS and filled_h1[c - 1] == title:
        sub_cols.append(c)
        c += 1
    blocks[col] = {
        "course_name": meta["name"].strip(),
        "course_code": meta["code"].strip(),
        "semester": int(meta["sem"]),
        "section": meta["sec"],
        "kind": meta["kind"],
        "cols": sub_cols,
        "subheaders": [header2[cc - 1] for cc in sub_cols],
    }
    col = c

print("Parsed blocks:")
for k, b in blocks.items():
    print(" ", b["course_code"], b["course_name"], "Sem", b["semester"], "Sec", b["section"], b["kind"], b["subheaders"])

# ---- Walk every student row and every block, emit long-format records ----
attendance_rows = []
assessment_rows = []
quality_log = []

def cellval(r, c):
    return ws.cell(row=r, column=c).value

def is_present(vals):
    return any(v is not None and v != "NA" for v in vals)

for r in range(3, NROWS + 1):
    usn = cellval(r, 1)
    name = cellval(r, 2)
    if usn is None:
        continue

    for start_col, b in blocks.items():
        raw_vals = [cellval(r, c) for c in b["cols"]]
        if not is_present(raw_vals):
            continue  # 'NA' block -> student not enrolled in this course/report

        rec = dict(zip(b["subheaders"], raw_vals))

        if b["kind"] == "Attendance":
            pct = rec.get("Attendance %")
            eng = rec.get("Classes Engaged")
            att = rec.get("Classes Attended")
            eligibility = None
            if isinstance(pct, (int, float)):
                if pct >= 75:
                    eligibility = "Eligible"
                elif pct >= 65:
                    eligibility = "Condonation"
                else:
                    eligibility = "Detained / NSAR"
            attendance_rows.append({
                "student_id": usn,
                "student_name": name,
                "course_code": b["course_code"],
                "course_name": b["course_name"],
                "semester": b["semester"],
                "section": b["section"],
                "classes_conducted": eng,
                "classes_attended": att,
                "attendance_pct": pct,
                "eligibility_status": eligibility,
                "duty_leave_credits": None,  # not present in source
            })

        else:  # IA Marks
            def parse_test(raw):
                """Return (numeric_value_or_None, flag)"""
                if raw == "A":
                    return 0.0, "ABSENT"
                if raw == "NE":
                    return None, "NOT_ELIGIBLE"
                if isinstance(raw, (int, float)):
                    return float(raw), None
                return None, None

            t1_raw = rec.get("TEST1")
            t2_raw = rec.get("TEST2")
            t3_raw = rec.get("TEST3")  # may be absent as a key for 2-test+ASN courses
            asn_raw = rec.get("ASN1")
            cie_raw = rec.get("CIE")

            t1, t1_flag = parse_test(t1_raw)
            t2, t2_flag = parse_test(t2_raw)
            t3, t3_flag = parse_test(t3_raw) if "TEST3" in rec else (None, None)
            asn, asn_flag = parse_test(asn_raw) if "ASN1" in rec else (None, None)

            flags = [f for f in (t1_flag, t2_flag, t3_flag, asn_flag) if f]
            if flags:
                quality_log.append({
                    "student_id": usn, "student_name": name,
                    "course_code": b["course_code"], "semester": b["semester"], "section": b["section"],
                    "issue": ",".join(sorted(set(flags))),
                })

            cie_reduced = float(cie_raw) if isinstance(cie_raw, (int, float)) else None
            cie_raw_equiv_100 = round(cie_reduced * 2, 2) if cie_reduced is not None else None

            assessment_rows.append({
                "student_id": usn,
                "student_name": name,
                "course_code": b["course_code"],
                "course_name": b["course_name"],
                "semester": b["semester"],
                "section": b["section"],
                "credits": None,          # not provided by source; fill from curriculum master
                "cie1": t1,
                "cie2": t2,
                "cie3": t3,
                "quiz1": None,            # this institution's IA scheme uses tests+assignment, not quizzes
                "quiz2": None,
                "el": asn,
                "cie_reduced": cie_reduced,      # authoritative, as computed & reported by the institution (out of 50)
                "cie_raw_equiv_100": cie_raw_equiv_100,  # derived x2, for schema compatibility ONLY (see notes)
                "see_raw": None,          # not present in source -- requires separate SEE mark sheet upload
                "see_reduced": None,
                "grand_total": None,
                "grade": None,
                "grade_points": None,
                "course_status": "in_progress",
                "assessment_notes": ";".join(flags) if flags else None,
            })

df_attendance = pd.DataFrame(attendance_rows)
df_assessments = pd.DataFrame(assessment_rows)
df_quality = pd.DataFrame(quality_log)

print("\nAttendance rows:", len(df_attendance))
print("Assessment rows:", len(df_assessments))
print("Quality flags:", len(df_quality))

df_attendance.to_csv("Attendance.csv", index=False)
df_assessments.to_csv("Course_Assessments.csv", index=False)
df_quality.to_csv("Data_Quality_Log.csv", index=False)

# ---- Build Students master (one row per USN) ----
USN_RE = re.compile(r"^\d[A-Z]{2}(\d{2})[A-Z]{2}\d{3}$")

def admission_year_from_usn(usn):
    m = USN_RE.match(usn)
    if not m:
        return None
    yy = int(m.group(1))
    # 2-digit year code; assume 2000s
    full = 2000 + yy
    return f"{full}-{str(full+1)[2:]}"

def batch_year_from_admission(adm_year_str):
    if not adm_year_str:
        return None
    start = int(adm_year_str.split("-")[0])
    end = start + 4
    return f"{start}-{str(end)[2:]}"

students = {}
for r in range(3, NROWS + 1):
    usn = cellval(r, 1)
    name = cellval(r, 2)
    if usn is None:
        continue
    students[usn] = {"student_id": usn, "name": name}

# derive latest known semester/section per student from assessment+attendance records
latest = {}
for rows_src in (attendance_rows, assessment_rows):
    for row in rows_src:
        sid = row["student_id"]
        sem = row["semester"]
        cur = latest.get(sid)
        if cur is None or sem > cur["semester"]:
            latest[sid] = {"semester": sem, "section": row["section"]}

students_rows = []
for usn, s in students.items():
    adm_year = admission_year_from_usn(usn)
    batch = batch_year_from_admission(adm_year)
    lt = latest.get(usn, {})
    students_rows.append({
        "student_id": usn,
        "name": s["name"],
        "email": None,            # TO_FILL from SIS / institutional email directory
        "phone": None,            # TO_FILL
        "department_id": "CSE",   # derived: all USNs carry the CS branch code
        "admission_year": adm_year,
        "admission_quota": None,  # TO_FILL from admissions office (KCET/COMEDK/Mgmt/SNQ/Lateral)
        "gender": None,           # TO_FILL
        "mentor_faculty_id": None,  # TO_FILL from proctor allotment sheet
        "semester": lt.get("semester"),          # inferred: latest semester seen in source reports
        "section": lt.get("section"),            # inferred: section at that latest semester
        "academic_term": None,    # TO_FILL - source doesn't carry an academic-year/term tag
        "batch_year": batch,
        "enrollment_status": None,  # TO_FILL from SIS (active/detained/discontinued/graduated)
    })

df_students = pd.DataFrame(students_rows).sort_values("student_id")
df_students.to_csv("Students.csv", index=False)

print("\nStudents master rows:", len(df_students))
print(df_students.head(3).to_string())
print("\nSample assessment rows:")
print(df_assessments.head(5).to_string())
print("\nSample attendance rows:")
print(df_attendance.head(5).to_string())
