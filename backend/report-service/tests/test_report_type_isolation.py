"""
Regression tests for Report Type Isolation.

Ensures:
1. General AI Report Builder no longer shares code path or prompts with SAR Criterion 4.
2. Generating reports of each report_type:
   - 'department_summary'
   - 'club_activity'
   - 'custom'
   - 'sar'
   asserts output does NOT contain SAR-specific headers (e.g. "4.1 Enrolment Ratio",
   "Placement Index", "Criterion 4") unless report_type is explicitly "sar".
3. For "custom", verifies LLM prompt contains:
   "Do not use NBA SAR Criterion 4 structure or terminology unless the user's instructions request it."
4. Verifies selective data fetching: only user-selected data sources are fetched.
5. Verifies general_report_builder never imports sar_generator.
"""

import sys
import os
import ast
import inspect
from unittest.mock import patch, MagicMock

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest
from report_service.general_report_builder import (
    ReportType,
    DataSource,
    build_general_report,
    fetch_selected_grounding_data,
    SYSTEM_PROMPT_CUSTOM,
)
from report_service.sar_generator import generate_sar_report
from render.docx_renderer import render_docx
from docx import Document
import io


SAR_SPECIFIC_HEADERS = [
    "4.1 Enrolment Ratio",
    "4.1",
    "Placement Index",
    "4.2 Success Rate",
    "4.2.1",
    "Success Rate without Backlogs",
    "4.2.2",
    "4.3",
    "4.4",
    "4.5",
    "4.6.1",
    "4.6.2",
    "4.6.3",
    "Criterion 4",
    "Criterion-4",
    "CRITERION-4",
    "Students Performance",
    "Students' Performance",
]


MOCK_CONFIG = {
    "ACADEMIC_DATA_SERVICE_URL": "http://mock-academic:8002",
    "NLP_RAG_SERVICE_URL": "http://mock-nlp:8005",
    "PREDICTION_SERVICE_URL": "http://mock-predict:8006",
    "REPORTS_DIR": "/tmp/reports",
}

MOCK_DEPT = {
    "id": 1,
    "code": "CSE",
    "name": "Computer Science & Engineering",
    "vision": "Excellence in computing",
    "mission": "Empower engineers",
}

MOCK_STUDENTS = [
    {"student_id": "STU001", "name": "Alice", "semester": 4, "attendance_pct": 85.0, "previous_gpa": 8.2, "final_result": "pass"},
    {"student_id": "STU002", "name": "Bob", "semester": 4, "attendance_pct": 78.0, "previous_gpa": 7.5, "final_result": "pass"},
]

MOCK_EVENTS = [
    {"id": 101, "title": "Hackathon 2025", "club_name": "Coding Club", "event_type": "Hackathon", "event_date": "2025-10-15", "attendee_count": 120},
    {"id": 102, "title": "AI Workshop", "club_name": "AI Society", "event_type": "Workshop", "event_date": "2025-11-20", "attendee_count": 80},
]

MOCK_FACULTY = [
    {"faculty_id": "FAC001", "name": "Dr. Sharma", "designation": "Professor", "qualification": "Ph.D.", "publications": 12, "research_projects": 2},
    {"faculty_id": "FAC002", "name": "Prof. Rao", "designation": "Assistant Professor", "qualification": "M.Tech", "publications": 3, "research_projects": 0},
]

MOCK_PLACEMENTS = {
    "years": [
        {"academic_year": "2024-25", "total_placed": 55, "higher_studies": 5, "entrepreneurs": 2, "placement_index": 0.85}
    ]
}


class TestReportTypeIsolation:

    def test_general_builder_does_not_import_sar_generator(self):
        """Invariant: general_report_builder must never import anything from sar_generator."""
        import report_service.general_report_builder as grb_module
        
        # 1. Module namespace check
        assert "sar_generator" not in grb_module.__dict__
        assert "build_sar_report_data" not in grb_module.__dict__
        assert "generate_sar_report" not in grb_module.__dict__

        # 2. AST source check
        src = inspect.getsource(grb_module)
        tree = ast.parse(src)
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for alias in node.names:
                    assert "sar_generator" not in alias.name, f"Forbidden import found: {alias.name}"
            elif isinstance(node, ast.ImportFrom):
                mod = node.module or ""
                assert "sar_generator" not in mod, f"Forbidden from-import found: {mod}"
                for alias in node.names:
                    assert "sar_generator" not in alias.name, f"Forbidden imported symbol: {alias.name}"

    @patch("data_client.fetch_department", return_value=MOCK_DEPT)
    @patch("data_client.fetch_all_students", return_value=MOCK_STUDENTS)
    @patch("data_client.fetch_approved_events", return_value=MOCK_EVENTS)
    @patch("data_client.fetch_all_faculty", return_value=MOCK_FACULTY)
    @patch("data_client.fetch_verified_placement_summary", return_value=MOCK_PLACEMENTS)
    @patch("llm_client.narrate")
    def test_reports_output_isolation_all_types(
        self, mock_narrate, mock_placements, mock_faculty, mock_events, mock_students, mock_dept
    ):
        """
        Generate one report of each report_type and assert that the output does NOT contain
        SAR-specific headers unless report_type is explicitly 'sar'.
        """
        mock_narrate.return_value = "Grounded narrative summary without any criteria references."

        report_types = ["department_summary", "club_activity", "custom", "sar"]

        for r_type in report_types:
            if r_type == "sar":
                # Mock SAR building components
                with patch("render.builder.build_report_data") as mock_sar_build:
                    from render.report_data import ReportData, ReportSection
                    sar_report = ReportData(
                        sar_format="ug_tier_ii_gapc_v4",
                        report_type="nba",
                        scope="full",
                        academic_year="2025-26",
                        generated_at="2025-10-01T00:00:00Z",
                        department=MOCK_DEPT,
                        sections=[
                            ReportSection(id="4", title="Students' Performance", marks=150, content_type="criterion_header", level=1),
                            ReportSection(id="4.1", title="Enrolment Ratio", marks=20, content_type="formula_table", level=2),
                            ReportSection(id="4.5", title="Placement Index", marks=40, content_type="formula_table", level=2),
                        ],
                    )
                    mock_sar_build.return_value = sar_report

                    res = generate_sar_report(
                        app_config=MOCK_CONFIG,
                        department_code="CSE",
                        academic_year="2025-26",
                    )
                    # For SAR, verify that SAR headers and branding ARE present
                    all_text = " ".join([s.title + " " + s.id for s in res.sections])
                    assert "4.1" in all_text and "Enrolment Ratio" in all_text
                    assert "4.5" in all_text and "Placement" in all_text

                    doc_sar_bytes = render_docx(res)
                    doc_sar = Document(io.BytesIO(doc_sar_bytes))
                    doc_sar_text = " ".join([p.text for p in doc_sar.paragraphs])
                    assert "NATIONAL BOARD OF ACCREDITATION" in doc_sar_text
                    assert "SELF-ASSESSMENT REPORT (SAR)" in doc_sar_text
                    assert "4.1" in doc_sar_text
                    assert "Enrolment Ratio" in doc_sar_text

            else:
                # Non-SAR general reports: department_summary, club_activity, custom
                res = build_general_report(
                    app_config=MOCK_CONFIG,
                    report_type=r_type,
                    department_id="CSE",
                    academic_year="2025-26",
                    selected_data=["student_records", "club_events", "faculty_data", "placement_data"],
                    section_titles=["Department Overview", "Initiatives", "Summary"] if r_type == "custom" else None,
                    instructions="Highlight academic success and student achievements" if r_type == "custom" else "",
                )

                # 1. Inspect section IDs and titles
                for sec in res.sections:
                    for sar_header in ["4.1 Enrolment Ratio", "Placement Index", "4.2.1", "4.6.1"]:
                        assert sar_header.lower() not in sec.title.lower(), (
                            f"SAR header '{sar_header}' leaked into section title '{sec.title}' in {r_type} report"
                        )
                        assert sar_header.lower() not in sec.id.lower(), (
                            f"SAR ID '{sar_header}' leaked into section id '{sec.id}' in {r_type} report"
                        )
                    assert sec.marks == 0, f"Non-SAR report should have 0 allocated marks, got {sec.marks}"
                    assert not sec.id.startswith("4."), f"Non-SAR section ID should not use '4.x' notation: {sec.id}"

                # 2. Render DOCX and verify no SAR branding / headers appear in generated document
                docx_bytes = render_docx(res)
                doc = Document(io.BytesIO(docx_bytes))
                doc_text = " ".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
                assert "NATIONAL BOARD OF ACCREDITATION" not in doc_text
                assert "SELF-ASSESSMENT REPORT (SAR)" not in doc_text
                assert "CRITERION-4" not in doc_text
                assert "4.1 Enrolment Ratio" not in doc_text
                assert "Placement Index" not in doc_text

                # 3. If jinja2 is installed, also verify rendered HTML
                try:
                    from render.pdf_renderer import _render_html
                    html_out = _render_html(res)
                    assert "National Board of Accreditation" not in html_out
                    assert "Self-Assessment Report (SAR)" not in html_out
                    assert "CRITERION-4" not in html_out
                    assert "4.1 Enrolment Ratio" not in html_out
                    assert "Placement Index" not in html_out
                except ImportError:
                    pass

    @patch("data_client.fetch_department", return_value=MOCK_DEPT)
    @patch("llm_client.narrate")
    def test_custom_report_prompt_includes_mandatory_directive(self, mock_narrate, mock_dept):
        """
        For 'custom' type, LLM prompt must explicitly state:
        'Do not use NBA SAR Criterion 4 structure or terminology unless the user's instructions request it.'
        """
        mock_narrate.return_value = "Custom grounded narrative."

        mandatory_directive = "Do not use NBA SAR Criterion 4 structure or terminology unless the user's instructions request it."
        assert mandatory_directive in SYSTEM_PROMPT_CUSTOM

        # Build custom report and check system_prompt passed to llm_client.narrate
        build_general_report(
            app_config=MOCK_CONFIG,
            report_type=ReportType.CUSTOM,
            department_id="CSE",
            academic_year="2025-26",
            selected_data=[],
            section_titles=["Intro Section", "Closing Notes"],
            instructions="Focus on first-year orientation.",
        )

        assert mock_narrate.called
        # Check system_prompt argument in mock_narrate call
        _, kwargs = mock_narrate.call_args
        called_prompt = kwargs.get("system_prompt", "")
        assert mandatory_directive in called_prompt
        assert "Focus on first-year orientation." in called_prompt

    @patch("data_client.fetch_department", return_value=MOCK_DEPT)
    @patch("data_client.fetch_all_students")
    @patch("data_client.fetch_approved_events")
    @patch("data_client.fetch_all_faculty")
    @patch("data_client.fetch_verified_placement_summary")
    def test_selective_data_fetching(
        self, mock_placements, mock_faculty, mock_events, mock_students, mock_dept
    ):
        """
        Verify that ONLY the user-selected data sources are fetched, rather than pulling
        the full SAR dataset by default.
        """
        # Case A: Only student_records checked
        mock_students.return_value = MOCK_STUDENTS
        raw, bullets = fetch_selected_grounding_data(
            academic_url="http://mock:8002",
            dept_code="CSE",
            academic_year="2025-26",
            selected_data=["student_records"],
        )
        assert mock_students.called
        assert not mock_events.called
        assert not mock_faculty.called
        assert not mock_placements.called
        assert "students" in raw
        assert "events" not in raw
        assert "faculty" not in raw

        # Reset mocks
        mock_students.reset_mock()
        mock_events.reset_mock()
        mock_faculty.reset_mock()
        mock_placements.reset_mock()

        # Case B: Only club_events checked
        mock_events.return_value = MOCK_EVENTS
        raw, bullets = fetch_selected_grounding_data(
            academic_url="http://mock:8002",
            dept_code="CSE",
            academic_year="2025-26",
            selected_data=["club_events"],
        )
        assert not mock_students.called
        assert mock_events.called
        assert not mock_faculty.called
        assert not mock_placements.called
        assert "events" in raw
        assert "students" not in raw

        # Reset mocks
        mock_students.reset_mock()
        mock_events.reset_mock()
        mock_faculty.reset_mock()
        mock_placements.reset_mock()

        # Case C: Only faculty_data checked
        mock_faculty.return_value = MOCK_FACULTY
        raw, bullets = fetch_selected_grounding_data(
            academic_url="http://mock:8002",
            dept_code="CSE",
            academic_year="2025-26",
            selected_data=["faculty_data"],
        )
        assert not mock_students.called
        assert not mock_events.called
        assert mock_faculty.called
        assert not mock_placements.called
        assert "faculty" in raw

        # Reset mocks
        mock_students.reset_mock()
        mock_events.reset_mock()
        mock_faculty.reset_mock()
        mock_placements.reset_mock()

        # Case D: Only placement_data checked
        mock_placements.return_value = MOCK_PLACEMENTS
        raw, bullets = fetch_selected_grounding_data(
            academic_url="http://mock:8002",
            dept_code="CSE",
            academic_year="2025-26",
            selected_data=["placement_data"],
        )
        assert not mock_students.called
        assert not mock_events.called
        assert not mock_faculty.called
        assert mock_placements.called
        assert "placements" in raw
